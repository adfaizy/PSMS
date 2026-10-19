"""
Register students from Excel onto youthgames.punjab.gov.pk
"""
from __future__ import annotations

import json
import re
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import requests
from PIL import Image

BASE = "https://youthgames.punjab.gov.pk"
FOLDER = Path(r"C:\Users\MOHID\Desktop\Photo")
XLSX = FOLDER / "Pubjab Youth Games.xlsx"
LOG = FOLDER / "registration_log.json"
PROGRESS = FOLDER / "registration_progress.json"

# Fixed IDs for this school batch
CATEGORY_ID = "1"  # Schools (Middle)
SUB_CATEGORY_ID = "1"  # Public
INSTITUTION_ID = "9152"  # GGES LUDHE KAY UCCHAY
DIVISION_ID = "6"  # Lahore
DISTRICT_ID = "21"  # Lahore
TEHSIL_ID = "76"  # Raiwind
UC_ID = "2346"  # UC No. 271, Jia Bagga

SPORT_IDS = {
    "Athletics": "1",
    "Cricket (Tape Ball)": "5",
}
KIT_IDS = {"Small": "1", "Medium": "2", "Large": "3", "XL": "4", "2XL": "5", "3XL": "6"}
JOGGER_IDS = {str(s): str(i) for i, s in enumerate(range(33, 47), start=1)}

MAX_PHOTO_BYTES = 250 * 1024  # under 256KB limit
REQUEST_DELAY = 2.5


def load_progress() -> dict:
    if PROGRESS.exists():
        return json.loads(PROGRESS.read_text(encoding="utf-8"))
    return {"done": {}, "failed": {}}


def save_progress(progress: dict) -> None:
    PROGRESS.write_text(json.dumps(progress, indent=2, ensure_ascii=False), encoding="utf-8")


def normalize_phone(phone: str) -> str:
    digits = re.sub(r"\D", "", str(phone or ""))
    if digits.startswith("92") and len(digits) == 12:
        digits = "0" + digits[2:]
    if len(digits) == 11 and digits.startswith("0"):
        return f"{digits[:4]}-{digits[4:]}"
    return str(phone)


def normalize_cnic(cnic: str) -> str:
    digits = re.sub(r"\D", "", str(cnic or ""))
    if len(digits) == 13:
        return f"{digits[:5]}-{digits[5:12]}-{digits[12:]}"
    return str(cnic)


def format_dob(dob) -> str:
    if isinstance(dob, datetime):
        return dob.strftime("%Y-%m-%d")
    return str(dob)[:10]


def prepare_photo(path: Path) -> tuple[BytesIO, str]:
    img = Image.open(path).convert("RGB")
    # Shrink until under size limit
    quality = 85
    w, h = img.size
    while True:
        buf = BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        if buf.tell() <= MAX_PHOTO_BYTES or (w < 200 and quality <= 40):
            buf.seek(0)
            return buf, path.name
        if buf.tell() > MAX_PHOTO_BYTES:
            if quality > 40:
                quality -= 10
            else:
                w = int(w * 0.8)
                h = int(h * 0.8)
                img = img.resize((max(w, 120), max(h, 120)), Image.Resampling.LANCZOS)


def solve_captcha(display: str) -> str:
    # e.g. "5 + 2 = ?" or "5 + 2"
    m = re.search(r"(\d+)\s*([+\-*/])\s*(\d+)", display)
    if not m:
        raise ValueError(f"Cannot parse captcha: {display}")
    a, op, b = int(m.group(1)), m.group(2), int(m.group(3))
    if op == "+":
        return str(a + b)
    if op == "-":
        return str(a - b)
    if op == "*":
        return str(a * b)
    if op == "/":
        return str(a // b)
    raise ValueError(f"Unknown op: {op}")


def new_session() -> requests.Session:
    s = requests.Session()
    s.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
            "Accept-Language": "en-US,en;q=0.9",
        }
    )
    return s


def get_csrf_and_captcha(session: requests.Session) -> tuple[str, str]:
    r = session.get(BASE + "/", timeout=60)
    r.raise_for_status()
    token = re.search(r'name="_token"\s+value="([^"]+)"', r.text)
    if not token:
        raise RuntimeError("CSRF token not found")
    c = session.get(
        BASE + f"/captcha-challenge/number?_={int(time.time()*1000)}",
        headers={"Accept": "application/json", "X-Requested-With": "XMLHttpRequest"},
        timeout=30,
    )
    c.raise_for_status()
    data = c.json()
    answer = solve_captcha(data.get("display") or data.get("question") or "")
    return token.group(1), answer


def register_one(session: requests.Session, row: dict) -> dict:
    token, captcha = get_csrf_and_captcha(session)
    photo_buf, photo_name = prepare_photo(row["photo_path"])

    data = {
        "_token": token,
        "disclaimer_accepted": "1",
        "category_id": CATEGORY_ID,
        "sub_category_id": SUB_CATEGORY_ID,
        "institution_id": INSTITUTION_ID,
        "institute_name": "",
        "first_name": row["full_name"],
        "last_name": row["father_name"],
        "cnic": row["cnic"],
        "gender": row["gender"],
        "date_of_birth": row["dob"],
        "division_id": DIVISION_ID,
        "district_id": DISTRICT_ID,
        "tehsil_id": TEHSIL_ID,
        "uc_id": UC_ID,
        "mobile_phone": row["phone"],
        "address": row["address"],
        "sport_id": row["sport_id"],
        "kit_size_id": row["kit_id"],
        "jogger_size_id": row["jogger_id"],
        "captcha": captcha,
    }
    files = {
        "profile_photo": (photo_name, photo_buf, "image/jpeg"),
    }
    resp = session.post(
        BASE + "/applicant",
        data=data,
        files=files,
        timeout=120,
        allow_redirects=True,
        headers={"Referer": BASE + "/#registration"},
    )
    text = resp.text
    # Persist raw response for debugging
    debug_path = FOLDER / f"last_response_{row['enrollment']}.html"
    debug_path.write_text(text, encoding="utf-8", errors="replace")

    success = False
    track_id = None
    error = None

    success_flag = 'data-registration-success="1"' in text
    server_errors = 'data-has-server-errors="1"' in text

    # Extract only registration form for field errors
    form_m = re.search(r'id="registration-form"[\s\S]*?</form>', text)
    form_html = form_m.group(0) if form_m else ""

    show_feedback = [
        re.sub(r"\s+", " ", m).strip()
        for m in re.findall(
            r'class="[^"]*show-feedback[^"]*"[^>]*>([^<]+)',
            form_html,
        )
        if m.strip()
    ]
    # CNIC already registered often appears in invalid-feedback near is-invalid cnic
    cnic_block = re.search(
        r'id="cnic"[\s\S]{0,500}?invalid-feedback[^>]*>([^<]+)',
        form_html,
        re.I,
    )
    cnic_msg = cnic_block.group(1).strip() if cnic_block else ""

    banner = ""
    bm = re.search(
        r'id="registration-errors"[\s\S]{0,1500}?</div>',
        text,
        re.I,
    )
    if bm:
        banner = re.sub(r"<[^>]+>", " ", bm.group(0))
        banner = re.sub(r"\s+", " ", banner).strip()

    # Track ID only from success areas (ignore placeholder CM-116816113)
    success_section = ""
    for sid in ("registration-success", "registration-success-banner", "registrationSuccessModal"):
        sm = re.search(rf'id="{sid}"[\s\S]{{0,4000}}', text)
        if sm:
            success_section += sm.group(0)
    tm = re.search(r"(CM-\d{6,})", success_section)
    if not tm and success_flag:
        # fallback: any CM id near success wording
        tm = re.search(
            r"(?:tracking|track|registration)\s*(?:id|number)?[^C]{0,40}(CM-\d{6,})",
            text,
            re.I,
        )
    if tm and tm.group(1) != "CM-116816113":
        track_id = tm.group(1)

    already = bool(
        re.search(r"already been registered|already registered", cnic_msg + " " + banner, re.I)
    )

    if success_flag:
        success = True
    elif already:
        error = "already_registered"
    elif server_errors:
        error = "validation: " + ("; ".join(show_feedback) or cnic_msg or banner or "server_errors")
    elif resp.status_code >= 400:
        error = f"http_{resp.status_code}"
    else:
        error = "unknown_response_not_marked_success"

    return {
        "ok": success,
        "already_registered": already,
        "status": resp.status_code,
        "track_id": track_id,
        "error": error,
        "url": str(resp.url),
        "banner": banner[:300],
        "show_feedback": show_feedback[:10],
        "cnic_msg": cnic_msg,
        "debug_html": str(debug_path),
    }


def read_students():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    students = []
    for r in range(2, ws.max_row + 1):
        enr = ws.cell(r, 21).value
        enr_key = str(enr).strip() if enr is not None else ""
        photo = FOLDER / f"{enr_key}.jpg"
        sport = ws.cell(r, 16).value
        kit = ws.cell(r, 17).value
        jog = str(ws.cell(r, 18).value)
        students.append(
            {
                "row": r,
                "enrollment": enr_key,
                "full_name": str(ws.cell(r, 5).value or "").strip(),
                "father_name": str(ws.cell(r, 6).value or "").strip(),
                "cnic": normalize_cnic(ws.cell(r, 7).value),
                "gender": str(ws.cell(r, 8).value or "").strip(),
                "dob": format_dob(ws.cell(r, 9).value),
                "phone": normalize_phone(ws.cell(r, 14).value),
                "address": str(ws.cell(r, 15).value or "").strip(),
                "sport_id": SPORT_IDS.get(str(sport).strip()),
                "kit_id": KIT_IDS.get(str(kit).strip()),
                "jogger_id": JOGGER_IDS.get(jog),
                "photo_path": photo if photo.exists() else None,
                "sport": sport,
                "kit": kit,
                "jogger": jog,
            }
        )
    return students


def main(limit: int | None = None, dry_run: bool = False):
    students = read_students()
    progress = load_progress()
    log = []

    with_photo = [s for s in students if s["photo_path"]]
    without = [s for s in students if not s["photo_path"]]
    print(f"Total: {len(students)} | with photo: {len(with_photo)} | without: {len(without)}")

    if dry_run:
        for s in with_photo[:3]:
            print("SAMPLE", {k: s[k] for k in s if k != "photo_path"}, "photo", s["photo_path"])
        return

    session = new_session()
    done_count = 0
    for s in with_photo:
        if limit is not None and done_count >= limit:
            break
        key = s["enrollment"] or s["cnic"]
        if key in progress.get("done", {}):
            print(f"SKIP already done {key}")
            continue
        if not s["sport_id"] or not s["kit_id"] or not s["jogger_id"]:
            msg = f"MISSING MAP for {key}: sport={s['sport']} kit={s['kit']} jog={s['jogger']}"
            print(msg)
            progress.setdefault("failed", {})[key] = msg
            save_progress(progress)
            continue

        print(f"Registering {key} {s['full_name']} ...")
        try:
            result = register_one(session, s)
        except Exception as e:
            result = {"ok": False, "error": str(e)}
            print("  ERROR", e)
        print(
            "  ->",
            {
                k: result.get(k)
                for k in ("ok", "already_registered", "status", "track_id", "error")
            },
        )
        log.append({"student": key, "name": s["full_name"], **{k: v for k, v in result.items() if k != "debug_html"}})
        if result.get("ok") or result.get("already_registered"):
            progress.setdefault("done", {})[key] = {
                "name": s["full_name"],
                "track_id": result.get("track_id"),
                "already_registered": bool(result.get("already_registered")),
                "at": datetime.now().isoformat(timespec="seconds"),
            }
            progress.get("failed", {}).pop(key, None)
        else:
            progress.setdefault("failed", {})[key] = {
                "name": s["full_name"],
                "error": result.get("error"),
                "show_feedback": result.get("show_feedback"),
                "cnic_msg": result.get("cnic_msg"),
            }
        save_progress(progress)
        LOG.write_text(json.dumps(log, indent=2, ensure_ascii=False), encoding="utf-8")
        done_count += 1
        time.sleep(REQUEST_DELAY)

    print("\nWithout photo (skipped):")
    for s in without:
        print(f"  {s['enrollment']} {s['full_name']}")
    print(f"\nDone so far: {len(progress.get('done', {}))} | Failed: {len(progress.get('failed', {}))}")


if __name__ == "__main__":
    import sys

    limit = 1
    dry = False
    if "--all" in sys.argv:
        limit = None
    if "--dry" in sys.argv:
        dry = True
    if "--limit" in sys.argv:
        limit = int(sys.argv[sys.argv.index("--limit") + 1])
    main(limit=limit, dry_run=dry)
